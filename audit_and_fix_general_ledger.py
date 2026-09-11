import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import re
import os
import sys
import collections

# -----------------------------------------------------------------------------
# CATEGORY COMPATIBILITY & EQUIVALENCE GROUPS
# -----------------------------------------------------------------------------

EQUIVALENCE_GROUPS = [
    # 1. Client Revenue Cycle (A/R, Services, Invoicing, Collections)
    {"accounts receivable (a/r)", "services", "accounts receivable", "services income", "income", "sales", "revenue"},
    
    # 2. Subcontractor / Direct Vendor Cycle (A/P, COGS - Vendor, Consulting Fees)
    {"accounts payable (a/p)", "cogs - vendor", "consulting fees", "accounts payable", "cost of goods sold", "vendor cogs", "cogs"},
    
    # 3. Owner Distributions / Draws
    {"owners distribution", "shareholders' equity:distributions", "distributions", "shareholders' equity", "owner's distribution"},
    
    # 4. Payroll Processing
    {"payroll expenses:payroll processing fee", "payroll processing fee", "payroll processing fees"},
    
    # 5. Payroll Taxes
    {"payroll wages and tax to pay:payroll tax to pay", "payroll tax to pay", "payroll expenses:payroll taxes", "payroll taxes"},
    
    # 6. Payroll Wages & Salaries
    {"payroll wages payable", "salaries & wages", "cogs - vendor:salaries & wages", "payroll wages"},
    
    # 7. Rent & Facilities
    {"rent:building & land rent", "building & land rent", "rent"},
]

def normalize_cat(cat):
    return str(cat).strip().lower()

def are_categories_compatible(cat1, cat2):
    """Check if two category strings are legitimate accounting counterparts or sub-accounts."""
    c1 = normalize_cat(cat1)
    c2 = normalize_cat(cat2)
    
    if c1 == c2:
        return True
        
    # Sub-account matching (e.g. "phone service" == "utilities:phone service")
    if c1.endswith(":" + c2) or c2.endswith(":" + c1):
        return True
    if ":" in c1 and ":" not in c2 and c1.split(":")[-1] == c2:
        return True
    if ":" in c2 and ":" not in c1 and c2.split(":")[-1] == c1:
        return True
        
    # Equivalence groups (e.g. Services vs Accounts Receivable (A/R))
    for grp in EQUIVALENCE_GROUPS:
        c1_in = any(c1 == m or c1.endswith(":" + m) or m.endswith(":" + c1) for m in grp)
        c2_in = any(c2 == m or c2.endswith(":" + m) or m.endswith(":" + c2) for m in grp)
        if c1_in and c2_in:
            return True
            
    return False

def clean_vendor(name, desc):
    """Normalize and extract merchant/vendor identity from Name or Description."""
    text = str(name).strip() if name and str(name).strip() and str(name).strip() != "None" else (str(desc).strip() if desc else "")
    if not text:
        return "Unknown"
    
    t_upper = text.upper()
    
    # Credit Cards
    if "AMERICAN EXPRESS" in t_upper or "AMEX" in t_upper:
        return "Amex CC"
    if "BANK OF AMERICA - CREDIT CARD" in t_upper or "BOA CC" in t_upper or "B OF A CREDIT CARD" in t_upper:
        return "BOA CC"
    # ADP Payroll
    if "ADP PAYROLL FEES" in t_upper or "ADP FEES" in t_upper:
        return "ADP Payroll Processing"
    if "ADP TAX" in t_upper:
        return "ADP Payroll Tax"
    if "ADP 401K" in t_upper:
        return "ADP 401k"
    if "ADP WAGE" in t_upper:
        return "ADP Wage Pay"
    # Food & Delivery
    if "DOORDASH" in t_upper or "DD *" in t_upper:
        if "TARGET" in t_upper or "ACE HARDWARE" in t_upper:
            return "DoorDash - Retail Supplies"
        return "DoorDash"
    if "UBER EATS" in t_upper:
        return "Uber Eats"
    if "UBER" in t_upper:
        return "Uber"
    # Tech & Software
    if "OPENAI" in t_upper or "CHATGPT" in t_upper:
        return "OpenAI"
    if "ANTHROPIC" in t_upper or "CLAUDE" in t_upper:
        return "Anthropic"
    if "LINKEDIN" in t_upper:
        return "LinkedIn"
    if "GOOGLE WORKSPACE" in t_upper or "GOOGLE" in t_upper:
        return "Google Workspace"
    if "MICROSOFT" in t_upper:
        return "Microsoft"
    if "AMAZON" in t_upper:
        return "Amazon"
    if "INTUIT" in t_upper or "QUICKBOOKS" in t_upper:
        return "Intuit QuickBooks"
    if "NETFLIX" in t_upper:
        return "Netflix"
    if "PRIME VIDEO" in t_upper:
        return "Prime Video"
    if "HP INSTANT INK" in t_upper:
        return "HP Instant Ink"
    # Facilities, Travel & Fuel
    if "BANNER PEST" in t_upper:
        return "Banner Pest Services"
    if "QUIK STOP" in t_upper:
        return "Quik Stop"
    if "UNION 76" in t_upper:
        return "Union 76"
    if "CHEVRON" in t_upper:
        return "Chevron"
    if "WAWA" in t_upper:
        return "Wawa"
    if "FASTRAK" in t_upper:
        return "Fastrak"
    if "SPOTHERO" in t_upper:
        return "SpotHero"
    # Telecom & Utilities
    if "DELMARVA" in t_upper:
        return "Delmarva Power"
    if "COMCAST" in t_upper or "XFINITY" in t_upper:
        return "Comcast / Xfinity"
    if "AT&T" in t_upper:
        return "AT&T"
    if "VERIZON" in t_upper:
        return "Verizon"
    if "RINGCENTRAL" in t_upper:
        return "RingCentral"
    if "VONAGE" in t_upper:
        return "Vonage"
    # Legal, Insurance & Subcontractors
    if "CSAA" in t_upper:
        return "CSAA Insurance"
    if "GLOBAL IMMIGRATION" in t_upper or "GLOBALIMMIGRATIO" in t_upper:
        return "Global Immigration Partners"
    if "SYNERGY BADMINTON" in t_upper:
        return "Synergy Badminton"
    if "FREEWHEEL BREWING" in t_upper:
        return "Freewheel Brewing"
    if "ACTIVATE PLANO" in t_upper:
        return "Activate Plano"
    if "CLUB 28 BADMINTON" in t_upper:
        return "Club 28 Badminton"
    if "ASTRAMIND" in t_upper:
        return "Astramind Solutions"
    if "ANERU" in t_upper:
        return "Aneru LLC"
    if "PTL CONSULTANCY" in t_upper:
        return "PTL Consultancy"
    if "SYSCONSULT" in t_upper:
        return "Sysconsult LLC"
    if "SAGE SOLUTIONS" in t_upper:
        return "Sage Solutions"
    if "AMERICAN TECHNOLOGY SERVICES" in t_upper:
        return "American Technology Services LLC"
    if "AGILISIUM" in t_upper:
        return "Agilisium Consulting LLC"
    if "MEDIX" in t_upper:
        return "Medix Staffing Solutions"
    if "NITYO" in t_upper:
        return "Nityo Infotech Corporation"
    if "INFINITE COMPUTER" in t_upper:
        return "Infinite Computer Solutions"
    if "PHOENIX STAFF" in t_upper:
        return "Phoenix Staff Inc"
    if "BIGBYTE" in t_upper:
        return "Bigbyte"
    if "SANA SOFTWARE" in t_upper:
        return "SANA Software Solutions"
    if "VIVID TECHNOLOGIES" in t_upper:
        return "Vivid Technologies Inc"
    if "VENDORPASS" in t_upper:
        return "Vendorpass Inc"
    if "RANDSTAND" in t_upper or "RANDSTAD" in t_upper:
        return "Randstad Digital"
    if "SANVI" in t_upper:
        return "Sanvi Consulting Services LLC"
    
    clean = re.sub(r"^(AplPay|TST\*\s*|DD\s*\*\s*|PY\s*\*\s*|PAYPAL\s*\*|SQ\s*\*)", "", text, flags=re.IGNORECASE).strip()
    return clean[:35]

def extract_historical_rules(hist_filepath):
    """Dynamically learn rules from historical workbook."""
    print(f"Loading Historical Data from {hist_filepath}...")
    hist_wb = openpyxl.load_workbook(hist_filepath, data_only=True)
    sheet = hist_wb.active
    rows = list(sheet.iter_rows(values_only=True))
    learned = collections.defaultdict(collections.Counter)
    
    for r in rows:
        c_dist = str(r[1]).strip() if r[1] is not None else ""
        c_name = str(r[5]).strip() if r[5] is not None else ""
        c_desc = str(r[6]).strip() if r[6] is not None else ""
        c_split = str(r[7]).strip() if r[7] is not None else ""
        
        if not c_dist or c_dist in ["Beginning Balance", "Distribution account"]:
            continue
            
        category = c_split if c_dist in ["Bank of America", "Amex CC", "BOA CC"] else c_dist
        vendor = clean_vendor(c_name, c_desc)
        
        if vendor and vendor != "Unknown" and category and "Uncategorized" not in category:
            learned[vendor][category] += 1
            
    rules = {}
    for v, cat_counts in learned.items():
        top_cat = cat_counts.most_common(1)[0][0]
        rules[v] = {
            "primary_category": top_cat,
            "occurrences": sum(cat_counts.values()),
            "all_categories": dict(cat_counts)
        }
    print(f"Successfully learned {len(rules)} historical vendor rules!")
    return rules

def audit_and_fix_general_ledger(curr_filepath, output_filepath, historical_rules):
    print(f"Auditing Current Ledger: {curr_filepath}...")
    wb = openpyxl.load_workbook(curr_filepath)
    sheet = wb["Sheet1"]
    
    header_bg = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    fill_cc_updated = PatternFill(start_color="D1ECF1", end_color="D1ECF1", fill_type="solid")
    fill_mismatch = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
    fill_uncategorized = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
    fill_matched = PatternFill(start_color="E8F8F5", end_color="E8F8F5", fill_type="solid")
    
    font_bold = Font(name="Arial", size=9, bold=True)
    font_regular = Font(name="Arial", size=9)
    font_danger = Font(name="Arial", size=9, color="721C24", bold=True)
    font_warning = Font(name="Arial", size=9, color="856404", bold=True)
    font_info = Font(name="Arial", size=9, color="0C5460", bold=True)
    font_success = Font(name="Arial", size=9, color="155724")
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    cell_k5 = sheet.cell(row=5, column=11, value="Audit Status")
    cell_l5 = sheet.cell(row=5, column=12, value="Audit Details & Suggested Category")
    for cell in [cell_k5, cell_l5]:
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = header_bg
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    discrepancy_list = []
    cc_updates_count = 0
    mismatch_count = 0
    uncat_count = 0
    matched_count = 0
    
    for r in range(6, sheet.max_row + 1):
        c_dist = sheet.cell(row=r, column=2).value
        c_date = sheet.cell(row=r, column=3).value
        c_type = sheet.cell(row=r, column=4).value
        c_num = sheet.cell(row=r, column=5).value
        cell_name = sheet.cell(row=r, column=6)
        c_name = cell_name.value
        c_desc = sheet.cell(row=r, column=7).value
        c_split = sheet.cell(row=r, column=8).value
        c_amt = sheet.cell(row=r, column=9).value
        
        if not c_dist or c_dist in ["Beginning Balance", "Distribution account"]:
            continue
            
        str_dist = str(c_dist).strip()
        str_type = str(c_type).strip() if c_type else ""
        str_name = str(c_name).strip() if c_name else ""
        str_desc = str(c_desc).strip() if c_desc else ""
        str_split = str(c_split).strip() if c_split else ""
        amt_val = float(c_amt) if c_amt is not None else 0.0
        
        # 1. Credit card payment vendor fix
        is_cc_payment = False
        target_cc_name = ""
        
        if str_type == "Credit Card Payment" or "CREDIT CARD Bill Payment" in str_desc or "AMERICAN EXPRESS DES:ACH PMT" in str_desc:
            is_cc_payment = True
            if "Amex CC" in str_dist or "Amex CC" in str_split or "AMERICAN EXPRESS" in str_desc:
                target_cc_name = "Amex CC"
            elif "BOA CC" in str_dist or "BOA CC" in str_split or "BANK OF AMERICA - CREDIT CARD" in str_desc:
                target_cc_name = "BOA CC"
            else:
                target_cc_name = "Credit Card"
        elif str_dist == "Bank of America" and str_split in ["Amex CC", "BOA CC"]:
            is_cc_payment = True
            target_cc_name = str_split
            
        if is_cc_payment:
            cell_name.value = target_cc_name
            status = "CC VENDOR UPDATED"
            note = f"Vendor Name auto-set to '{target_cc_name}'"
            row_fill = fill_cc_updated
            row_font = font_info
            cc_updates_count += 1
            discrepancy_list.append({
                "row": r, "date": c_date, "account": str_dist, "type": str_type,
                "name": target_cc_name, "desc": str_desc, "split": str_split,
                "amt": amt_val, "status": status, "issue": "Missing CC Vendor Name",
                "action": f"Set Vendor Name to '{target_cc_name}'"
            })
        else:
            vendor_key = clean_vendor(str_name, str_desc)
            
            if "Uncategorized" in str_dist or "Uncategorized" in str_split:
                status = "UNCATEGORIZED"
                note = "Uncategorized transaction requires reconciliation / invoice mapping."
                row_fill = fill_uncategorized
                row_font = font_danger
                uncat_count += 1
                discrepancy_list.append({
                    "row": r, "date": c_date, "account": str_dist, "type": str_type,
                    "name": str_name or vendor_key, "desc": str_desc, "split": str_split,
                    "amt": amt_val, "status": status, "issue": "Uncategorized Transaction",
                    "action": "Assign to specific client invoice or expense account"
                })
            elif vendor_key in historical_rules:
                rule_info = historical_rules[vendor_key]
                expected_category = rule_info["primary_category"]
                curr_category = str_split if str_dist in ["Bank of America", "Amex CC", "BOA CC"] else str_dist
                
                # Check compatibility with parent accounts and equivalence groups
                if not are_categories_compatible(curr_category, expected_category):
                    status = "CATEGORY MISMATCH"
                    note = f"Expected '{expected_category}' based on historical data, found '{curr_category}'."
                    row_fill = fill_mismatch
                    row_font = font_warning
                    mismatch_count += 1
                    discrepancy_list.append({
                        "row": r, "date": c_date, "account": str_dist, "type": str_type,
                        "name": str_name or vendor_key, "desc": str_desc, "split": str_split,
                        "amt": amt_val, "status": status, "issue": f"Mismatched Category (Found: {curr_category})",
                        "action": f"Reclassify to historical category: '{expected_category}'"
                    })
                else:
                    status = "MATCHED"
                    note = f"Verified: {expected_category}"
                    row_fill = fill_matched
                    row_font = font_success
                    matched_count += 1
            else:
                status = "MATCHED"
                note = "Standard ledger entry"
                row_fill = fill_matched
                row_font = font_success
                matched_count += 1
                
        cell_k = sheet.cell(row=r, column=11, value=status)
        cell_l = sheet.cell(row=r, column=12, value=note)
        cell_k.font = row_font
        cell_l.font = font_regular
        
        for col_idx in range(2, 13):
            sheet.cell(row=r, column=col_idx).fill = row_fill
            
    if "Audit & Discrepancies" in wb.sheetnames:
        del wb["Audit & Discrepancies"]
    audit_ws = wb.create_sheet(title="Audit & Discrepancies", index=0)
    audit_ws.views.sheetView[0].showGridLines = True
    
    audit_ws.merge_cells("A1:H1")
    title_cell = audit_ws["A1"]
    title_cell.value = "General Ledger Audit & Compliance Summary"
    title_cell.font = Font(name="Arial", size=14, bold=True, color="1F4E79")
    
    audit_ws["A2"] = f"Total Transactions Audited: {len(discrepancy_list) + matched_count} | Issues / Updates: {len(discrepancy_list)}"
    audit_ws["A2"].font = Font(name="Arial", size=10, italic=True, color="595959")
    
    kpis = [
        ("Total Discrepancies / Flags", len(discrepancy_list), "FFF3CD", "856404"),
        ("CC Vendor Payments Fixed", cc_updates_count, "D1ECF1", "0C5460"),
        ("Category Inconsistencies", mismatch_count, "FFF3CD", "856404"),
        ("Uncategorized Inflows/Outflows", uncat_count, "F8D7DA", "721C24"),
    ]
    
    col_start = 1
    for title, val, bg_hex, text_hex in kpis:
        c1 = audit_ws.cell(row=4, column=col_start, value=title)
        c2 = audit_ws.cell(row=5, column=col_start, value=val)
        c1.font = Font(name="Arial", size=9, bold=True, color=text_hex)
        c2.font = Font(name="Arial", size=16, bold=True, color=text_hex)
        c1.fill = PatternFill(start_color=bg_hex, end_color=bg_hex, fill_type="solid")
        c2.fill = PatternFill(start_color=bg_hex, end_color=bg_hex, fill_type="solid")
        c1.alignment = Alignment(horizontal="center")
        c2.alignment = Alignment(horizontal="center")
        audit_ws.merge_cells(start_row=4, start_column=col_start, end_row=4, end_column=col_start+1)
        audit_ws.merge_cells(start_row=5, start_column=col_start, end_row=5, end_column=col_start+1)
        col_start += 2
        
    headers = ["Row #", "Date", "Account", "Type", "Vendor / Name", "Description", "Amount", "Status", "Identified Issue", "Recommended Action"]
    for col_idx, h_text in enumerate(headers, 1):
        cell = audit_ws.cell(row=7, column=col_idx, value=h_text)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = header_bg
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    for row_idx, item in enumerate(discrepancy_list, 8):
        audit_ws.cell(row=row_idx, column=1, value=item["row"]).alignment = Alignment(horizontal="center")
        audit_ws.cell(row=row_idx, column=2, value=item["date"]).alignment = Alignment(horizontal="center")
        audit_ws.cell(row=row_idx, column=3, value=item["account"])
        audit_ws.cell(row=row_idx, column=4, value=item["type"])
        audit_ws.cell(row=row_idx, column=5, value=item["name"]).font = font_bold
        audit_ws.cell(row=row_idx, column=6, value=item["desc"])
        c_amt = audit_ws.cell(row=row_idx, column=7, value=item["amt"])
        c_amt.number_format = "$#,##0.00;($#,##0.00);\"-\""
        
        c_stat = audit_ws.cell(row=row_idx, column=8, value=item["status"])
        if item["status"] == "UNCATEGORIZED":
            c_stat.fill = fill_uncategorized
            c_stat.font = font_danger
        elif item["status"] == "CATEGORY MISMATCH":
            c_stat.fill = fill_mismatch
            c_stat.font = font_warning
        else:
            c_stat.fill = fill_cc_updated
            c_stat.font = font_info
        c_stat.alignment = Alignment(horizontal="center")
        
        audit_ws.cell(row=row_idx, column=9, value=item["issue"])
        audit_ws.cell(row=row_idx, column=10, value=item["action"]).font = font_bold
        
        for c in range(1, 11):
            audit_ws.cell(row=row_idx, column=c).border = thin_border
            
    for ws in [audit_ws, sheet]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                if cell.row < 4 and ws == audit_ws:
                    continue
                if '\n' in val_str:
                    val_str = max(val_str.split('\n'), key=len)
                max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)
            
    wb.save(output_filepath)
    print(f"Audit complete! Saved output to: {output_filepath}")
    print(f" - Evaluated: {matched_count + len(discrepancy_list)} rows (100% styled)")
    print(f" - Credit Card Payments Fixed: {cc_updates_count}")
    print(f" - Category Mismatches vs History: {mismatch_count}")
    print(f" - Uncategorized Items: {uncat_count}")
    print(f" - Matched Transactions: {matched_count}")

if __name__ == "__main__":
    hist_file = "VIRTUSSOLUTIONS LLC_General Ledger (_Historical data.xlsx"
    curr_file = "VIRTUSSOLUTIONS LLC_General Ledger (1).xlsx"
    out_file = "VIRTUSSOLUTIONS_LLC_Audited_General_Ledger.xlsx"
    
    rules = extract_historical_rules(hist_file)
    audit_and_fix_general_ledger(curr_file, out_file, rules)
